import io
import json
import re
import uuid
from datetime import datetime, timezone

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn
from sqlalchemy import text

from app.agent.quick_models import reasoning_for
from app.config import Settings, get_settings
from app.db.session import SessionLocal
from app.repositories import conversation_repo
from app.services import conversation_export
from app.services.conversation_export import (
    ExportHeader,
    ExportLine,
    ExportTranslationError,
    bangla_digits,
    needs_translation,
    render_docx,
    translate,
)


@pytest.mark.parametrize(
    "text,target,expected",
    [
        ("What is my balance?", "bn", True),
        ("amar balance koto", "bn", True),
        ("আমার ব্যালেন্স কত?", "bn", False),
        ("আমার balance কত?", "bn", True),
        ("*121# 500", "bn", False),
        ("What is my balance?", "en", True),
        ("amar balance koto", "en", True),
        ("আমার ব্যালেন্স কত?", "en", True),
        ("*121# 500", "en", False),
    ],
)
def test_needs_translation(text, target, expected):
    assert needs_translation(text, target) is expected


def _llm(handler):
    """An httpx client whose OpenRouter calls go to `handler(items) -> reply content`."""
    requests: list[dict] = []

    def respond(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        requests.append(body)
        items = json.loads(body["messages"][-1]["content"])
        content = handler(items)
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]})

    return httpx.Client(transport=httpx.MockTransport(respond)), requests


def _echo_upper(items):
    return json.dumps([{"i": item["i"], "text": item["text"].upper()} for item in items])


def test_translate_sends_only_needed_texts():
    client, requests = _llm(_echo_upper)
    texts = ["hello", "আমার ব্যালেন্স কত?", "*121#", "amar net slow"]
    out = translate(texts, "bn", client=client, settings=get_settings())
    assert out == ["HELLO", "আমার ব্যালেন্স কত?", "*121#", "AMAR NET SLOW"]
    sent = [item["text"] for item in json.loads(requests[0]["messages"][-1]["content"])]
    assert sent == ["hello", "amar net slow"]


def test_translate_banglish_both_ways():
    for target in ("bn", "en"):
        client, requests = _llm(_echo_upper)
        translate(["amar net slow keno"], target, client=client, settings=get_settings())
        sent = json.loads(requests[0]["messages"][-1]["content"])
        assert [item["text"] for item in sent] == ["amar net slow keno"]


def test_translate_nothing_needed_makes_no_call():
    client, requests = _llm(_echo_upper)
    assert translate(["আমি ভালো আছি"], "bn", client=client, settings=get_settings()) == [
        "আমি ভালো আছি"
    ]
    assert requests == []


def test_translate_chunks_above_40():
    client, requests = _llm(_echo_upper)
    out = translate([f"line {n}" for n in range(41)], "bn", client=client, settings=get_settings())
    assert len(requests) == 2
    assert out[40] == "LINE 40"


def test_translate_accepts_fenced_json():
    client, _ = _llm(lambda items: "```json\n" + _echo_upper(items) + "\n```")
    assert translate(["hi"], "bn", client=client, settings=get_settings()) == ["HI"]


def test_translate_missing_index_fails():
    client, _ = _llm(lambda items: json.dumps([{"i": items[0]["i"], "text": "x"}]))
    with pytest.raises(ExportTranslationError):
        translate(["one", "two"], "bn", client=client, settings=get_settings())


def test_translate_bad_json_fails():
    client, _ = _llm(lambda items: "not json")
    with pytest.raises(ExportTranslationError):
        translate(["one"], "bn", client=client, settings=get_settings())


def test_translate_http_error_fails():
    server_error = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(500)))
    with pytest.raises(ExportTranslationError):
        translate(["one"], "bn", client=server_error, settings=get_settings())

    def time_out(request):
        raise httpx.ReadTimeout("slow", request=request)

    timeout = httpx.Client(transport=httpx.MockTransport(time_out))
    with pytest.raises(ExportTranslationError):
        translate(["one"], "bn", client=timeout, settings=get_settings())


def test_translate_prompt_rules():
    client, requests = _llm(_echo_upper)
    translate(["hi"], "bn", client=client, settings=get_settings())
    body = requests[0]
    system = body["messages"][0]["content"]
    # The bench showed loanwords left in Latin letters ("আমার balance কত?") and
    # Bangla digits kept in English output ("১২০ taka").
    for rule in ("Banglish", "telco", "প্যাকেজ", "*121#", "ব্যালেন্স", "digits 0-9"):
        assert rule in system
    assert body["model"] == get_settings().export_translate_model
    # The code default (a local .env may pick a cheaper model while testing).
    assert Settings.model_fields["export_translate_model"].default == "google/gemini-3.5-flash"
    # Gemini 3.5 rejects reasoning "none" ("Reasoning is mandatory"); the shared
    # per-model setting is minimal for a model it doesn't list.
    assert body["reasoning"] == reasoning_for(body["model"])
    assert body["reasoning"] != {"effort": "none"}


# 2026-10-30 08:32 UTC is 14:32 in Dhaka.
_STARTED = datetime(2026, 10, 30, 8, 32, tzinfo=timezone.utc)


def _header(**overrides) -> ExportHeader:
    fields = dict(
        started_at=_STARTED,
        ended_at=_STARTED,
        channel="inbound_chat",
        is_phone_call=False,
        customer_name="Rahim Uddin",
        customer_msisdn="01712345678",
    )
    return ExportHeader(**(fields | overrides))


_LINES = [
    ExportLine(role="assistant", text="Hello! How can I help?", at=_STARTED),
    ExportLine(role="user", text="Call 01712345678 please", at=_STARTED),
]


def _text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_render_en_labels_and_title():
    text = _text(render_docx(_header(), _LINES, "en"))
    assert "telco chat transcript" in text
    assert "Assistant:" in text and "User:" in text
    assert "Hello! How can I help?" in text
    assert "14:32" in text and "Web chat" in text
    assert "telco." not in text


def test_render_bn_labels_and_digits():
    text = _text(render_docx(_header(), _LINES, "bn"))
    assert "telco চ্যাট ট্রান্সক্রিপ্ট" in text
    assert "সহকারী:" in text and "গ্রাহক:" in text
    assert "৩০ অক্টোবর ২০২৬" in text
    assert "Call 01712345678 please" in text
    assert "Assistant:" not in text


def test_render_header_customer():
    with_customer = _text(render_docx(_header(), _LINES, "en"))
    assert "Rahim Uddin" in with_customer and "01712345678" in with_customer
    guest = _text(render_docx(_header(customer_name=None, customer_msisdn=None), _LINES, "en"))
    assert "Customer" not in guest


def test_render_sets_cs_font():
    doc = Document(io.BytesIO(render_docx(_header(), _LINES, "bn")))
    runs = [run for p in doc.paragraphs for run in p.runs]
    assert runs
    for run in runs:
        assert run._element.rPr.rFonts.get(qn("w:cs")) == "Nirmala UI"


def test_bangla_digits():
    assert bangla_digits("2026-10-30 14:32") == "২০২৬-১০-৩০ ১৪:৩২"


# --- export + routes -------------------------------------------------------

EXPORT_ROOM_PREFIX = "pytest-export-"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def stored():
    """Factory for stored conversations, deleted afterwards."""
    yield _store
    db = SessionLocal()
    try:
        db.execute(
            text("delete from conversations where room_name like :p"),
            {"p": EXPORT_ROOM_PREFIX + "%"},
        )
        db.commit()
    finally:
        db.close()


def _store(
    visitor_id: uuid.UUID | None = None,
    *,
    ended: bool = True,
    is_phone_call: bool = False,
    texts: tuple[tuple[str, str | None], ...] = (
        ("assistant", "Hello!"),
        ("user", "amar balance koto"),
    ),
) -> uuid.UUID:
    db = SessionLocal()
    try:
        conversation = conversation_repo.create(
            db,
            channel="inbound_chat",
            room_name=f"{EXPORT_ROOM_PREFIX}{uuid.uuid4().hex[:8]}",
            is_phone_call=is_phone_call,
            visitor_id=visitor_id,
        )
        now = datetime.now(timezone.utc)
        for seq, (role, body) in enumerate(texts, start=1):
            conversation_repo.add_message(
                db, conversation.id, seq=seq, role=role, text=body, created_at=now
            )
        conversation_repo.add_message(
            db,
            conversation.id,
            seq=len(texts) + 1,
            role="tool",
            text="balance 500",
            tool_name="get_balance",
            tool_args={},
            created_at=now,
        )
        if ended:
            conversation_repo.end(db, conversation.id)
        return conversation.id
    finally:
        db.close()


@pytest.fixture
def fake_translate(monkeypatch):
    """translate() replaced: records its input, upper-cases what it would send."""
    calls: list[list[str]] = []

    def fake(texts, target, *, client, settings):
        calls.append(list(texts))
        return [t.upper() if needs_translation(t, target) else t for t in texts]

    monkeypatch.setattr(conversation_export, "translate", fake)
    return calls


def test_export_skips_empty_messages(stored, fake_translate):
    conversation_id = stored(texts=(("assistant", "Hi"), ("user", None), ("user", "   ")))
    db = SessionLocal()
    try:
        conversation = conversation_repo.get(db, conversation_id)
        name, data = conversation_export.export_conversation(db, conversation, "bn")
    finally:
        db.close()
    assert fake_translate == [["Hi"]]
    body = _text(data)
    assert "HI" in body and "balance 500" not in body
    assert name.endswith("-bn.docx")


def test_admin_export_ok(api, stored, fake_translate):
    conversation_id = stored()
    res = api.get(f"/api/conversations/{conversation_id}/export", params={"lang": "bn"})
    assert res.status_code == 200
    assert res.headers["content-type"] == DOCX_TYPE
    assert re.search(
        r'filename="telco-chat-\d{4}-\d{2}-\d{2}-\d{4}-bn\.docx"',
        res.headers["content-disposition"],
    )
    body = _text(res.content)
    assert "AMAR BALANCE KOTO" in body and "balance 500" not in body


def test_admin_export_requires_admin(anon_api, stored):
    conversation_id = stored()
    res = anon_api.get(f"/api/conversations/{conversation_id}/export", params={"lang": "en"})
    assert res.status_code == 401


def test_admin_export_active_409(api, stored, fake_translate):
    conversation_id = stored(ended=False)
    res = api.get(f"/api/conversations/{conversation_id}/export", params={"lang": "en"})
    assert res.status_code == 409


def test_admin_export_translation_502(api, stored, monkeypatch):
    def fail(*args, **kwargs):
        raise ExportTranslationError("down")

    monkeypatch.setattr(conversation_export, "translate", fail)
    conversation_id = stored()
    res = api.get(f"/api/conversations/{conversation_id}/export", params={"lang": "en"})
    assert res.status_code == 502
    assert res.json() == {"detail": "translation failed"}


def test_admin_export_bad_lang_422(api, stored):
    conversation_id = stored()
    res = api.get(f"/api/conversations/{conversation_id}/export", params={"lang": "fr"})
    assert res.status_code == 422


def test_admin_export_missing_404(api):
    res = api.get(f"/api/conversations/{uuid.uuid4()}/export", params={"lang": "en"})
    assert res.status_code == 404


def _visitor_get(client, conversation_id, visitor_id, lang="en"):
    return client.get(
        f"/api/public/conversations/{conversation_id}/export",
        params={"lang": lang},
        headers={"X-Visitor-Id": str(visitor_id)},
    )


def test_visitor_export_ok(anon_api, stored, fake_translate):
    visitor = uuid.uuid4()
    res = _visitor_get(anon_api, stored(visitor), visitor)
    assert res.status_code == 200
    assert res.headers["content-type"] == DOCX_TYPE
    assert res.headers["cache-control"] == "no-store"


def test_visitor_export_other_visitor_404(anon_api, stored, fake_translate):
    res = _visitor_get(anon_api, stored(uuid.uuid4()), uuid.uuid4())
    assert res.status_code == 404


def test_visitor_export_phone_call_404(anon_api, stored, fake_translate):
    visitor = uuid.uuid4()
    res = _visitor_get(anon_api, stored(visitor, is_phone_call=True), visitor)
    assert res.status_code == 404


def test_visitor_export_active_409(anon_api, stored, fake_translate):
    visitor = uuid.uuid4()
    res = _visitor_get(anon_api, stored(visitor, ended=False), visitor)
    assert res.status_code == 409


def test_cors_exposes_the_filename(api):
    origin = get_settings().cors_origins[0]
    res = api.get(f"/api/conversations/{uuid.uuid4()}/export?lang=en", headers={"Origin": origin})
    assert "content-disposition" in res.headers.get("access-control-expose-headers", "").lower()


def test_admin_detail_returns_reply_language(api, stored):
    conversation_id = stored()
    db = SessionLocal()
    try:
        db.execute(
            text("update conversations set reply_language = 'bn' where id = :i"),
            {"i": conversation_id},
        )
        db.commit()
    finally:
        db.close()
    assert api.get(f"/api/conversations/{conversation_id}").json()["reply_language"] == "bn"
