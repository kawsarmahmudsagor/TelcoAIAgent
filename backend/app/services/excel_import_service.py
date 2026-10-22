"""Imports dummy CRM leads from an .xlsx file into the `clients` table.

Expected header row (order doesn't matter, matched by name, case-insensitive):
full_name, phone, email, lead_source, budget_min, budget_max, preferred_location,
preferred_property_type, bedrooms_wanted, bathrooms_wanted, move_in_date, rent_or_buy, notes

Only `full_name` is required per row; everything else is optional.
See backend/app/data/seed_sample.xlsx for a matching sample file.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.client import RentOrBuy
from app.repositories import client_repo

_KNOWN_COLUMNS = {
    "full_name",
    "phone",
    "email",
    "lead_source",
    "budget_min",
    "budget_max",
    "preferred_location",
    "preferred_property_type",
    "bedrooms_wanted",
    "bathrooms_wanted",
    "move_in_date",
    "rent_or_buy",
    "notes",
}


def _coerce(column: str, value) -> object:
    if value is None or value == "":
        return None
    if column in {"budget_min", "budget_max", "bathrooms_wanted"}:
        return float(value)
    if column == "bedrooms_wanted":
        return int(value)
    if column == "move_in_date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return date.fromisoformat(str(value))
    if column == "rent_or_buy":
        try:
            return RentOrBuy(str(value).strip().lower())
        except ValueError:
            return RentOrBuy.undecided
    return str(value).strip()


def import_clients(db: Session, file_bytes: bytes) -> dict:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active

    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]

    imported = 0
    skipped = 0
    errors: list[str] = []

    for row_index, raw_row in enumerate(rows, start=2):
        row = dict(zip(header, raw_row, strict=False))
        full_name = row.get("full_name")
        if not full_name:
            skipped += 1
            errors.append(f"row {row_index}: missing full_name, skipped")
            continue

        try:
            data = {
                column: _coerce(column, row.get(column))
                for column in _KNOWN_COLUMNS
                if column in row
            }
            data["full_name"] = str(full_name).strip()
            client_repo.create(db, data)
            imported += 1
        except (ValueError, TypeError) as exc:
            skipped += 1
            errors.append(f"row {row_index}: {exc}")

    return {"imported": imported, "skipped": skipped, "errors": errors}
