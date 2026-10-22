"""Chunk + embed files and text into a Chroma collection (ported from Sellora's
rag.ingestion, sync).

Every chunk carries the caller's metadata (rag_source_id at least), which is
what lets retrieval filter by source and lets a source's vectors be deleted
without knowing their doc ids. Bookkeeping keys are excluded from both the
embedded text and the LLM-facing text.

Re-ingesting under the same doc_id replaces the previous chunks (delete, then
insert), so a reindex never accumulates duplicates.
"""

from __future__ import annotations

import json
import logging
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from llama_index.core import SimpleDirectoryReader
from llama_index.core.ingestion.pipeline import run_transformations
from llama_index.core.schema import Document
from openpyxl import load_workbook

from app.rag.index import get_index, transformations

logger = logging.getLogger("rag.ingestion")

_TRANSIENT_MARKERS = ("429", "500", "502", "503", "504", "UNAVAILABLE", "RESOURCE_EXHAUSTED")

# Header-row keywords locating the question/answer columns of a Query/Reply sheet.
_QUERY_HEADER_HINTS = ("quer", "question")
_REPLY_HEADER_HINTS = ("repl", "answ")

# Chroma only stores scalar metadata values.
_SCALAR_TYPES = (str, int, float, bool)


@dataclass
class EmbedTarget:
    """Where and how to embed: the collection, the model that writes it, the
    key for that model, and the chunking."""

    collection: str
    embedding_model: str
    api_key: str
    chunk_size: int = 512
    chunk_overlap: int = 50


@dataclass
class TextDoc:
    doc_id: str
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


def _retry_on_transient_error(func, max_retries: int = 3, initial_delay: float = 1.0):
    for attempt in range(1, max_retries + 1):
        try:
            return func()
        except Exception as exc:
            transient = any(marker in str(exc) for marker in _TRANSIENT_MARKERS)
            if transient and attempt < max_retries:
                delay = initial_delay * (2 ** (attempt - 1))
                logger.warning(
                    "embedding call failed (attempt %d/%d), retrying in %.1fs: %s",
                    attempt, max_retries, delay, exc,
                )  # fmt: skip
                time.sleep(delay)
            else:
                raise


def _clean_metadata(metadata: dict[str, Any] | None) -> dict[str, Any]:
    cleaned: dict[str, Any] = {}
    for key, value in (metadata or {}).items():
        if value is None:
            continue
        cleaned[key] = value if isinstance(value, _SCALAR_TYPES) else str(value)
    return cleaned


def _prepare(document: Document, doc_id: str, metadata: dict[str, Any] | None) -> Document:
    extra = _clean_metadata(metadata)
    document.id_ = doc_id
    document.metadata = {**document.metadata, **extra}
    hidden = list(extra)
    document.excluded_embed_metadata_keys = list({*document.excluded_embed_metadata_keys, *hidden})
    document.excluded_llm_metadata_keys = list({*document.excluded_llm_metadata_keys, *hidden})
    return document


def embed_documents(documents: list[Document], target: EmbedTarget) -> int:
    """Chunks and embeds already-prepared documents (each with its doc id and
    metadata), replacing whatever was stored under those doc ids. Returns the
    number of chunks embedded."""
    index = get_index(target.collection, target.embedding_model, target.api_key)
    for doc_id in {document.id_ for document in documents}:
        index.vector_store.delete(doc_id)

    nodes = run_transformations(
        documents, transformations(target.chunk_size, target.chunk_overlap)
    )
    if not nodes:
        logger.warning("no chunks to embed for %d document(s)", len(documents))
        return 0
    _retry_on_transient_error(lambda: index.insert_nodes(nodes))
    logger.info("embedded %d chunk(s) into %s", len(nodes), target.collection)
    return len(nodes)


def _find_column(header_cells: list[str], hints: tuple[str, ...]) -> int | None:
    for position, cell in enumerate(header_cells):
        lowered = cell.lower()
        if any(hint in lowered for hint in hints):
            return position
    return None


def read_qa_spreadsheet(path: Path) -> str:
    """Every worksheet of a Query/Reply spreadsheet as `### question` + answer
    blocks, so the markdown parser makes one chunk per pair. Columns come from
    the header row; without a recognisable header, A and B are used. Rows with
    a blank question or answer are skipped."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
            query_col = _find_column(header, _QUERY_HEADER_HINTS)
            reply_col = _find_column(header, _REPLY_HEADER_HINTS)
            if query_col is not None and reply_col is not None:
                data_rows = rows[1:]
            else:
                query_col, reply_col = 0, 1
                data_rows = rows
            for row in data_rows:
                if len(row) <= max(query_col, reply_col):
                    continue
                query = str(row[query_col]).strip() if row[query_col] is not None else ""
                reply = str(row[reply_col]).strip() if row[reply_col] is not None else ""
                if query and reply:
                    blocks.append(f"### {query}\n\n{reply}\n")
    finally:
        workbook.close()
    return "\n".join(blocks)


def read_qa_json(path: Path) -> str:
    """A JSON file as markdown blocks: a list of {question, answer} (or
    query/reply, q/a) objects becomes one `### question` section per pair;
    other shapes are kept as headed JSON."""
    data = json.loads(path.read_text(encoding="utf-8"))
    blocks: list[str] = []
    if isinstance(data, list):
        for item in data:
            if isinstance(item, dict):
                question = item.get("question") or item.get("query") or item.get("q")
                answer = item.get("answer") or item.get("reply") or item.get("a")
                if question and answer:
                    blocks.append(f"### {str(question).strip()}\n\n{str(answer).strip()}\n")
                else:
                    dumped = json.dumps(item, ensure_ascii=False, indent=2)
                    blocks.append(f"```json\n{dumped}\n```\n")
            else:
                blocks.append(str(item).strip())
    elif isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, (dict, list)):
                dumped = json.dumps(value, ensure_ascii=False, indent=2)
                blocks.append(f"### {key}\n\n```json\n{dumped}\n```\n")
            else:
                blocks.append(f"### {key}\n\n{str(value).strip()}\n")
    else:
        blocks.append(str(data))
    return "\n".join(blocks)


def load_file_documents(path: Path) -> list[Document]:
    """md/txt/pdf via LlamaIndex's file readers; xlsx/json as Q&A markdown."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        text = read_qa_spreadsheet(path)
        if not text.strip():
            raise ValueError(f"no question/answer rows found in {path.name}")
        return [Document(text=text)]
    if suffix == ".json":
        text = read_qa_json(path)
        if not text.strip():
            raise ValueError(f"no readable JSON content in {path.name}")
        return [Document(text=text)]
    documents = SimpleDirectoryReader(input_files=[str(path)]).load_data()
    documents = [document for document in documents if document.text.strip()]
    if not documents:
        raise ValueError(f"no readable text in {path.name}")
    return documents


def ingest_file(
    path: Path, doc_id: str, target: EmbedTarget, *, metadata: dict[str, Any] | None = None
) -> int:
    """Chunk + embed a file under `doc_id`. Returns the chunk count."""
    documents = [_prepare(document, doc_id, metadata) for document in load_file_documents(path)]
    return embed_documents(documents, target)


def ingest_text(
    text: str, doc_id: str, target: EmbedTarget, *, metadata: dict[str, Any] | None = None
) -> int:
    """Chunk + embed a block of text under `doc_id`. Returns the chunk count."""
    if not text or not text.strip():
        raise ValueError(f"refusing to embed empty text for doc_id={doc_id}")
    return embed_documents([_prepare(Document(text=text), doc_id, metadata)], target)


def ingest_texts(docs: list[TextDoc], target: EmbedTarget) -> int:
    """Embeds several texts (each its own doc id) in one batch."""
    documents = [
        _prepare(Document(text=doc.text), doc.doc_id, doc.metadata)
        for doc in docs
        if doc.text.strip()
    ]
    if not documents:
        return 0
    return embed_documents(documents, target)
